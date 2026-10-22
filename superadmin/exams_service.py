import csv
import io
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.db.models import (
    Count,
    Exists,
    F,
    OuterRef,
    Q,
    QuerySet,
    Subquery,
    Sum,
)
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime

from admins.models import AdminInstitution
from api.models import Exam, Question, StudentAttempt
from api.utils.schedule import get_exam_schedule_phase
from superadmin.reports_service import sanitize_formula_cell

User = get_user_model()

ALLOWED_STATUSES = {'generating', 'draft', 'failed', 'published', 'closed'}
ALLOWED_PHASES = {'upcoming', 'live', 'ended'}
ALLOWED_QUESTION_TYPES = {'mcq', 'true_false', 'fill_blanks', 'short', 'long'}
ALLOWED_SORTS = {
    'created': '-created_at',
    'created_asc': 'created_at',
    'title': 'title',
    'title_desc': '-title',
    'start': 'starts_at',
    'start_desc': '-starts_at',
    'end': 'ends_at',
    'end_desc': '-ends_at',
    'marks': '-total_marks',
    'marks_asc': 'total_marks',
    'shortfall': '-marks_shortfall',
    'attempts': '-attempts_count',
}


def filter_exams_queryset(
    filters: Dict[str, Any],
    base_qs: Optional[QuerySet] = None,
) -> Tuple[QuerySet, Dict[str, Any]]:
    """
    Applies comprehensive server-side filtering on the Exam queryset.
    Returns (filtered_queryset, applied_filters_summary).
    Raises ValidationError on malformed or illegal parameter values.
    """
    now = timezone.now()
    applied: Dict[str, Any] = {}

    qs = base_qs if base_qs is not None else Exam.objects.all()

    # Optimized select_related and annotations to prevent N+1 queries
    qs = qs.select_related(
        'institution',
        'created_by',
        'created_by__faculty_profile',
    ).annotate(
        attempts_count=Count('attempts', distinct=True),
        pending_grading_count=Count(
            'attempts',
            filter=Q(attempts__status__in=[StudentAttempt.Status.SUBMITTED, StudentAttempt.Status.AUTO_SUBMITTED]),
            distinct=True,
        ),
        flagged_students_count=Count(
            'attempts',
            filter=Q(attempts__violation_count__gt=0),
            distinct=True,
        ),
        questions_count=Count('questions', distinct=True),
    )

    # 1. Search Query: title, subject, department, faculty username/email, or exact exam id
    search = str(filters.get('search') or filters.get('q') or '').strip()
    if search:
        search_q = (
            Q(title__icontains=search)
            | Q(subject__icontains=search)
            | Q(department__icontains=search)
            | Q(created_by__username__icontains=search)
            | Q(created_by__email__icontains=search)
            | Q(created_by__first_name__icontains=search)
            | Q(created_by__last_name__icontains=search)
            | Q(institution__institution_name__icontains=search)
            | Q(institution__institution_code__icontains=search)
        )
        if search.isdigit():
            search_q |= Q(id=int(search))
        qs = qs.filter(search_q)
        applied['search'] = search

    # 2. Institution Filter
    institution_id = filters.get('institution_id')
    if institution_id not in (None, '', 'all'):
        try:
            inst_id_int = int(institution_id)
            qs = qs.filter(institution_id=inst_id_int)
            applied['institution_id'] = inst_id_int
        except (ValueError, TypeError):
            raise ValidationError({'institution_id': 'Invalid institution ID format.'})

    # 3. Department Filter
    department = str(filters.get('department', '')).strip()
    if department and department != 'all':
        qs = qs.filter(department__iexact=department)
        applied['department'] = department

    # 4. Faculty Filter
    faculty_id = filters.get('faculty_id')
    if faculty_id not in (None, '', 'all'):
        try:
            fac_id_int = int(faculty_id)
            qs = qs.filter(created_by_id=fac_id_int)
            applied['faculty_id'] = fac_id_int
        except (ValueError, TypeError):
            raise ValidationError({'faculty_id': 'Invalid faculty user ID format.'})

    # 5. Multi-Select Status Filter
    raw_status = filters.get('status')
    if raw_status:
        statuses: List[str] = []
        if isinstance(raw_status, list):
            statuses = [str(s).strip().lower() for s in raw_status if str(s).strip()]
        elif isinstance(raw_status, str):
            statuses = [s.strip().lower() for s in raw_status.split(',') if s.strip()]

        # Filter out 'all'
        valid_requested = [s for s in statuses if s != 'all']
        if valid_requested:
            invalid = [s for s in valid_requested if s not in ALLOWED_STATUSES]
            if invalid:
                raise ValidationError({'status': f'Invalid status(es): {", ".join(invalid)}'})
            qs = qs.filter(status__in=valid_requested)
            applied['status'] = valid_requested

    # 6. Schedule Phase Filter (upcoming, live, ended)
    raw_phase = filters.get('phase')
    if raw_phase:
        phases: List[str] = []
        if isinstance(raw_phase, list):
            phases = [str(p).strip().lower() for p in raw_phase if str(p).strip()]
        elif isinstance(raw_phase, str):
            phases = [p.strip().lower() for p in raw_phase.split(',') if p.strip()]

        valid_phases = [p for p in phases if p in ALLOWED_PHASES]
        if valid_phases:
            phase_q = Q()
            if 'upcoming' in valid_phases:
                phase_q |= Q(status=Exam.Status.PUBLISHED, starts_at__gt=now)
            if 'live' in valid_phases:
                phase_q |= Q(
                    status=Exam.Status.PUBLISHED,
                    starts_at__lte=now,
                ) & (Q(ends_at__isnull=True) | Q(ends_at__gt=now))
            if 'ended' in valid_phases:
                phase_q |= Q(status=Exam.Status.CLOSED) | Q(
                    status=Exam.Status.PUBLISHED,
                    ends_at__lte=now,
                )
            qs = qs.filter(phase_q)
            applied['phase'] = valid_phases

    # 7. Schedule Time Windows
    # Opens From / To
    opens_from = filters.get('opens_from')
    if opens_from:
        d = parse_datetime(str(opens_from)) or (parse_date(str(opens_from)) and timezone.make_aware(datetime.combine(parse_date(str(opens_from)), datetime.min.time())))
        if not d:
            raise ValidationError({'opens_from': 'Invalid date format (ISO-8601 expected).'})
        qs = qs.filter(starts_at__gte=d)
        applied['opens_from'] = d.isoformat()

    opens_to = filters.get('opens_to')
    if opens_to:
        d = parse_datetime(str(opens_to)) or (parse_date(str(opens_to)) and timezone.make_aware(datetime.combine(parse_date(str(opens_to)), datetime.max.time())))
        if not d:
            raise ValidationError({'opens_to': 'Invalid date format (ISO-8601 expected).'})
        qs = qs.filter(starts_at__lte=d)
        applied['opens_to'] = d.isoformat()

    # Closes From / To
    closes_from = filters.get('closes_from')
    if closes_from:
        d = parse_datetime(str(closes_from)) or (parse_date(str(closes_from)) and timezone.make_aware(datetime.combine(parse_date(str(closes_from)), datetime.min.time())))
        if not d:
            raise ValidationError({'closes_from': 'Invalid date format (ISO-8601 expected).'})
        qs = qs.filter(ends_at__gte=d)
        applied['closes_from'] = d.isoformat()

    closes_to = filters.get('closes_to')
    if closes_to:
        d = parse_datetime(str(closes_to)) or (parse_date(str(closes_to)) and timezone.make_aware(datetime.combine(parse_date(str(closes_to)), datetime.max.time())))
        if not d:
            raise ValidationError({'closes_to': 'Invalid date format (ISO-8601 expected).'})
        qs = qs.filter(ends_at__lte=d)
        applied['closes_to'] = d.isoformat()

    # Closing within next 24 hours
    closing_soon = filters.get('closing_soon')
    if closing_soon in (True, 'true', '1', 1):
        in_24h = now + timedelta(hours=24)
        qs = qs.filter(
            status=Exam.Status.PUBLISHED,
            ends_at__gt=now,
            ends_at__lte=in_24h,
        )
        applied['closing_soon'] = True

    # Missing Schedule (Drafts without configured starts_at or ends_at)
    missing_schedule = filters.get('missing_schedule')
    if missing_schedule in (True, 'true', '1', 1):
        qs = qs.filter(Q(starts_at__isnull=True) | Q(ends_at__isnull=True))
        applied['missing_schedule'] = True

    # 8. Created Date Range
    created_from = filters.get('created_from')
    if created_from:
        d = parse_datetime(str(created_from)) or (parse_date(str(created_from)) and timezone.make_aware(datetime.combine(parse_date(str(created_from)), datetime.min.time())))
        if not d:
            raise ValidationError({'created_from': 'Invalid date format (ISO-8601 expected).'})
        qs = qs.filter(created_at__gte=d)
        applied['created_from'] = d.isoformat()

    created_to = filters.get('created_to')
    if created_to:
        d = parse_datetime(str(created_to)) or (parse_date(str(created_to)) and timezone.make_aware(datetime.combine(parse_date(str(created_to)), datetime.max.time())))
        if not d:
            raise ValidationError({'created_to': 'Invalid date format (ISO-8601 expected).'})
        qs = qs.filter(created_at__lte=d)
        applied['created_to'] = d.isoformat()

    # 9. Marks & Shortfall
    target_marks_min = filters.get('target_marks_min') or filters.get('min_marks')
    if target_marks_min not in (None, ''):
        try:
            val = int(target_marks_min)
            qs = qs.filter(total_marks__gte=val)
            applied['target_marks_min'] = val
        except ValueError:
            raise ValidationError({'target_marks_min': 'Must be an integer.'})

    target_marks_max = filters.get('target_marks_max') or filters.get('max_marks')
    if target_marks_max not in (None, ''):
        try:
            val = int(target_marks_max)
            qs = qs.filter(total_marks__lte=val)
            applied['target_marks_max'] = val
        except ValueError:
            raise ValidationError({'target_marks_max': 'Must be an integer.'})

    has_shortfall = filters.get('has_shortfall')
    if has_shortfall in (True, 'true', '1', 1):
        qs = qs.filter(marks_shortfall__gt=0)
        applied['has_shortfall'] = True
    elif has_shortfall in (False, 'false', '0', 0):
        qs = qs.filter(marks_shortfall=0)
        applied['has_shortfall'] = False

    shortfall_min = filters.get('shortfall_min')
    if shortfall_min not in (None, ''):
        try:
            val = int(shortfall_min)
            qs = qs.filter(marks_shortfall__gte=val)
            applied['shortfall_min'] = val
        except ValueError:
            raise ValidationError({'shortfall_min': 'Must be an integer.'})

    shortfall_max = filters.get('shortfall_max')
    if shortfall_max not in (None, ''):
        try:
            val = int(shortfall_max)
            qs = qs.filter(marks_shortfall__lte=val)
            applied['shortfall_max'] = val
        except ValueError:
            raise ValidationError({'shortfall_max': 'Must be an integer.'})

    # 10. Duration & Question Count
    duration_min = filters.get('duration_min')
    if duration_min not in (None, ''):
        try:
            val = int(duration_min)
            qs = qs.filter(duration_minutes__gte=val)
            applied['duration_min'] = val
        except ValueError:
            raise ValidationError({'duration_min': 'Must be an integer.'})

    duration_max = filters.get('duration_max')
    if duration_max not in (None, ''):
        try:
            val = int(duration_max)
            qs = qs.filter(duration_minutes__lte=val)
            applied['duration_max'] = val
        except ValueError:
            raise ValidationError({'duration_max': 'Must be an integer.'})

    questions_min = filters.get('questions_count_min')
    if questions_min not in (None, ''):
        try:
            val = int(questions_min)
            qs = qs.filter(questions_count__gte=val)
            applied['questions_count_min'] = val
        except ValueError:
            raise ValidationError({'questions_count_min': 'Must be an integer.'})

    questions_max = filters.get('questions_count_max')
    if questions_max not in (None, ''):
        try:
            val = int(questions_max)
            qs = qs.filter(questions_count__lte=val)
            applied['questions_count_max'] = val
        except ValueError:
            raise ValidationError({'questions_count_max': 'Must be an integer.'})

    # 11. Question Types Present
    raw_qtypes = filters.get('question_types')
    if raw_qtypes:
        qtypes: List[str] = []
        if isinstance(raw_qtypes, list):
            qtypes = [str(q).strip().lower() for q in raw_qtypes if str(q).strip()]
        elif isinstance(raw_qtypes, str):
            qtypes = [q.strip().lower() for q in raw_qtypes.split(',') if q.strip()]

        valid_qtypes = [q for q in qtypes if q in ALLOWED_QUESTION_TYPES]
        if valid_qtypes:
            qs = qs.filter(questions__question_type__in=valid_qtypes).distinct()
            applied['question_types'] = valid_qtypes

    # 12. Attempts & Proctoring Flags
    has_attempts = filters.get('has_attempts')
    if has_attempts in (True, 'true', '1', 1):
        qs = qs.filter(attempts_count__gt=0)
        applied['has_attempts'] = True
    elif has_attempts in (False, 'false', '0', 0):
        qs = qs.filter(attempts_count=0)
        applied['has_attempts'] = False

    has_pending_grading = filters.get('has_pending_grading')
    if has_pending_grading in (True, 'true', '1', 1):
        qs = qs.filter(pending_grading_count__gt=0)
        applied['has_pending_grading'] = True

    has_flagged_students = filters.get('has_flagged_students')
    if has_flagged_students in (True, 'true', '1', 1):
        qs = qs.filter(flagged_students_count__gt=0)
        applied['has_flagged_students'] = True

    # 13. AI Model Used & Generation Source
    ai_model = str(filters.get('ai_model', '')).strip()
    if ai_model and ai_model != 'all':
        qs = qs.filter(ai_model_used__icontains=ai_model)
        applied['ai_model'] = ai_model

    source = str(filters.get('source', '')).strip().lower()
    if source in ('ai_generated', 'manual'):
        if source == 'ai_generated':
            qs = qs.filter(Q(ai_model_used__isnull=False) & ~Q(ai_model_used=''))
        else:
            qs = qs.filter(Q(ai_model_used__isnull=True) | Q(ai_model_used=''))
        applied['source'] = source

    # 14. Sorting
    sort_key = str(filters.get('sort_by', 'created')).strip()
    db_sort = ALLOWED_SORTS.get(sort_key, '-created_at')
    # Stable tiebreaker on id
    qs = qs.order_by(db_sort, '-id')
    applied['sort_by'] = sort_key

    return qs, applied


def serialize_exam_row(exam: Exam) -> Dict[str, Any]:
    """Serializes an Exam model instance into a structured dictionary."""
    now = timezone.now()
    fac_name = ''
    fac_email = ''
    if exam.created_by:
        fac_name = f'{exam.created_by.first_name} {exam.created_by.last_name}'.strip() or exam.created_by.username
        fac_email = exam.created_by.email

    closing_soon = bool(
        exam.status == Exam.Status.PUBLISHED
        and exam.ends_at
        and now < exam.ends_at <= now + timedelta(hours=24)
    )

    return {
        'id': exam.id,
        'title': exam.title,
        'subject': exam.subject or '-',
        'department': exam.department or 'General',
        'institution': {
            'id': exam.institution.id if exam.institution else None,
            'name': exam.institution.institution_name if exam.institution else 'Unassigned',
            'code': exam.institution.institution_code if exam.institution else '-',
        },
        'faculty': {
            'id': exam.created_by.id if exam.created_by else None,
            'username': exam.created_by.username if exam.created_by else 'Unknown',
            'full_name': fac_name,
            'email': fac_email,
        },
        'status': exam.status,
        'schedule_phase': get_exam_schedule_phase(exam, now=now),
        'duration_minutes': exam.duration_minutes,
        'total_marks': exam.total_marks,
        'target_total_marks': exam.target_total_marks,
        'marks_shortfall': exam.marks_shortfall,
        'starts_at': exam.starts_at.isoformat() if exam.starts_at else None,
        'ends_at': exam.ends_at.isoformat() if exam.ends_at else None,
        'closing_soon': closing_soon,
        'is_immutable': exam.is_immutable,
        'ai_model_used': exam.ai_model_used or 'Manual',
        'attempts_count': getattr(exam, 'attempts_count', 0),
        'pending_grading_count': getattr(exam, 'pending_grading_count', 0),
        'flagged_students_count': getattr(exam, 'flagged_students_count', 0),
        'questions_count': getattr(exam, 'questions_count', 0),
        'created_at': exam.created_at.isoformat() if exam.created_at else None,
    }


def get_exams_filter_metadata(institution_id: Optional[int] = None) -> Dict[str, Any]:
    """Returns available choices and options for dependent filter dropdowns."""
    institutions = list(
        AdminInstitution.objects.all()
        .order_by('institution_name')
        .values('id', 'institution_name', 'institution_code')
    )

    # Departments
    dept_qs = Exam.objects.exclude(department='').values_list('department', flat=True).distinct()
    if institution_id:
        dept_qs = dept_qs.filter(institution_id=institution_id)
    departments = sorted(list(set(dept_qs)))

    # Faculty list
    faculty_qs = User.objects.filter(created_exams__isnull=False).distinct()
    if institution_id:
        faculty_qs = faculty_qs.filter(created_exams__institution_id=institution_id)
    faculty = []
    for u in faculty_qs.select_related('faculty_profile').order_by('username')[:200]:
        name = f'{u.first_name} {u.last_name}'.strip() or u.username
        dept = u.faculty_profile.department if hasattr(u, 'faculty_profile') else ''
        faculty.append({'id': u.id, 'name': name, 'email': u.email, 'department': dept})

    # AI Models used
    ai_models = sorted(
        list(
            set(
                Exam.objects.exclude(ai_model_used='')
                .values_list('ai_model_used', flat=True)
                .distinct()
            )
        )
    )

    # Status counts platform-wide
    status_counts = {}
    for st in ALLOWED_STATUSES:
        status_counts[st] = Exam.objects.filter(status=st).count()

    return {
        'institutions': institutions,
        'departments': departments,
        'faculty': faculty,
        'ai_models': ai_models,
        'status_counts': status_counts,
        'allowed_statuses': list(ALLOWED_STATUSES),
        'allowed_phases': list(ALLOWED_PHASES),
        'allowed_question_types': list(ALLOWED_QUESTION_TYPES),
    }


def export_exams_to_csv_string(exams_qs: QuerySet) -> str:
    """Exports filtered exams to sanitized CSV string."""
    output = io.StringIO()
    writer = csv.writer(output)

    headers = [
        'ID',
        'Title',
        'Subject',
        'Department',
        'Institution',
        'Institution Code',
        'Faculty',
        'Faculty Email',
        'Status',
        'Schedule Phase',
        'Duration (Min)',
        'Total Marks',
        'Target Marks',
        'Shortfall',
        'Starts At (UTC)',
        'Ends At (UTC)',
        'Attempts Count',
        'Pending Grading',
        'Flagged Students',
        'AI Model Used',
        'Created At (UTC)',
    ]
    writer.writerow([sanitize_formula_cell(h) for h in headers])

    for e in exams_qs:
        row_data = serialize_exam_row(e)
        writer.writerow([
            sanitize_formula_cell(row_data['id']),
            sanitize_formula_cell(row_data['title']),
            sanitize_formula_cell(row_data['subject']),
            sanitize_formula_cell(row_data['department']),
            sanitize_formula_cell(row_data['institution']['name']),
            sanitize_formula_cell(row_data['institution']['code']),
            sanitize_formula_cell(row_data['faculty']['full_name']),
            sanitize_formula_cell(row_data['faculty']['email']),
            sanitize_formula_cell(row_data['status']),
            sanitize_formula_cell(row_data['schedule_phase']),
            sanitize_formula_cell(row_data['duration_minutes']),
            sanitize_formula_cell(row_data['total_marks']),
            sanitize_formula_cell(row_data['target_total_marks'] or '-'),
            sanitize_formula_cell(row_data['marks_shortfall']),
            sanitize_formula_cell(row_data['starts_at'] or '-'),
            sanitize_formula_cell(row_data['ends_at'] or '-'),
            sanitize_formula_cell(row_data['attempts_count']),
            sanitize_formula_cell(row_data['pending_grading_count']),
            sanitize_formula_cell(row_data['flagged_students_count']),
            sanitize_formula_cell(row_data['ai_model_used']),
            sanitize_formula_cell(row_data['created_at']),
        ])

    return output.getvalue()


def export_exams_to_xlsx_bytes(exams_qs: QuerySet) -> bytes:
    """Exports filtered exams to stylized Excel (.xlsx) binary."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Global Exams'

    header_font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    data_font = Font(name='Calibri', size=9, color='334155')
    border_side = Side(style='thin', color='E2E8F0')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    headers = [
        'ID',
        'Title',
        'Subject',
        'Department',
        'Institution',
        'Faculty',
        'Status',
        'Phase',
        'Duration',
        'Marks',
        'Shortfall',
        'Starts At',
        'Ends At',
        'Attempts',
        'Flagged',
        'AI Model',
    ]

    ws.append([f'ExamAI Global Exams Export — Generated {timezone.now().strftime("%Y-%m-%d %H:%M:%S UTC")}'])
    ws.cell(row=1, column=1).font = Font(name='Calibri', size=13, bold=True, color='1E293B')
    ws.append([])

    ws.append([sanitize_formula_cell(h) for h in headers])
    header_row = 3
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_idx not in (2, 5, 6) else 'left')

    current_row = 4
    for e in exams_qs:
        row_data = serialize_exam_row(e)
        row_vals = [
            sanitize_formula_cell(row_data['id']),
            sanitize_formula_cell(row_data['title']),
            sanitize_formula_cell(row_data['subject']),
            sanitize_formula_cell(row_data['department']),
            sanitize_formula_cell(row_data['institution']['name']),
            sanitize_formula_cell(row_data['faculty']['full_name']),
            sanitize_formula_cell(row_data['status']),
            sanitize_formula_cell(row_data['schedule_phase']),
            sanitize_formula_cell(f"{row_data['duration_minutes']}m"),
            sanitize_formula_cell(f"{row_data['total_marks']}/{row_data['target_total_marks'] or '-'}"),
            sanitize_formula_cell(row_data['marks_shortfall']),
            sanitize_formula_cell(row_data['starts_at'] or '-'),
            sanitize_formula_cell(row_data['ends_at'] or '-'),
            sanitize_formula_cell(row_data['attempts_count']),
            sanitize_formula_cell(row_data['flagged_students_count']),
            sanitize_formula_cell(row_data['ai_model_used']),
        ]
        ws.append(row_vals)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.font = data_font
            c.border = cell_border
        current_row += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(40, max(12, max_len + 3))

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
