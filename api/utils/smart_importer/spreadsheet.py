"""
Smart Spreadsheet Parser for Question Bank Import (CSV / XLSX).
Handles title rows, multi-encoding, delimiter sniffing, formula injection protection,
and automatic or user-defined column mapping.
"""

import csv
import io
import re
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from .synonyms import (
    COLUMN_TARGETS,
    TARGET_CORRECT_ANSWER,
    TARGET_IGNORE,
    TARGET_OPTION_A,
    TARGET_OPTION_B,
    TARGET_OPTION_C,
    TARGET_OPTION_D,
    TARGET_OPTION_E,
    TARGET_OPTION_F,
    TARGET_OPTIONS,
    TARGET_QUESTION_TEXT,
    match_column_name,
    normalize_header_string,
)

DANGEROUS_FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def sanitize_cell_value(val: Any) -> Any:
    """Neutralize Excel/CSV formula injection and strip trailing/leading whitespace."""
    if val is None:
        return ''
    if isinstance(val, (int, float, Decimal)):
        return val
    s = str(val).strip()
    while s and s[0] in DANGEROUS_FORMULA_PREFIXES:
        s = s[1:].strip()
    return s


def detect_csv_delimiter(sample_text: str) -> str:
    """Sniffs the most likely CSV delimiter from sample text."""
    delimiters = [',', ';', '\t', '|']
    sniffer = csv.Sniffer()
    try:
        dialect = sniffer.sniff(sample_text[:4096], delimiters=delimiters)
        return dialect.delimiter
    except Exception:
        # Fallback: count occurrences in non-empty lines
        lines = [line for line in sample_text.splitlines() if line.strip()][:10]
        if not lines:
            return ','
        counts = {d: sum(line.count(d) for line in lines) for d in delimiters}
        best = max(counts.items(), key=lambda x: x[1])
        return best[0] if best[1] > 0 else ','


def detect_header_row(raw_grid: List[List[Any]], max_scan_rows: int = 15) -> int:
    """
    Scans the first N rows of a spreadsheet grid to detect the true header row.
    Finds the row with the highest number of recognizable question bank column synonyms.
    Returns: 0-based index of the detected header row. Defaults to 0.
    """
    if not raw_grid:
        return 0

    best_row_idx = 0
    best_match_score = -1.0

    for idx, row in enumerate(raw_grid[:max_scan_rows]):
        if not row or not any(str(c).strip() for c in row if c is not None):
            continue

        non_empty = [c for c in row if c is not None and str(c).strip()]
        if not non_empty:
            continue

        # Count how many cells in this row match a valid column synonym
        matched_targets = set()
        total_score = 0.0

        for cell in non_empty:
            cell_str = str(cell).strip()
            target, score = match_column_name(cell_str, threshold=0.75)
            if target != TARGET_IGNORE and score > 0.6:
                matched_targets.add(target)
                total_score += score

        # Prioritize rows containing the essential anchors: question_text or options or correct_answer
        anchor_bonus = 0.0
        if TARGET_QUESTION_TEXT in matched_targets:
            anchor_bonus += 2.0
        if TARGET_CORRECT_ANSWER in matched_targets:
            anchor_bonus += 1.5
        if TARGET_OPTION_A in matched_targets or TARGET_OPTIONS in matched_targets:
            anchor_bonus += 1.5

        composite_score = total_score + anchor_bonus
        if composite_score > best_match_score and len(matched_targets) >= 1:
            best_match_score = composite_score
            best_row_idx = idx

    return best_row_idx


def parse_csv_grid(file_bytes: bytes) -> Tuple[List[List[Any]], str, str]:
    """
    Decodes CSV bytes with encoding detection and delimiter sniffing.
    Returns: (raw_grid, detected_encoding, detected_delimiter)
    """
    detected_encoding = 'utf-8'
    text = ""

    for enc in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            text = file_bytes.decode(enc)
            detected_encoding = enc
            break
        except UnicodeDecodeError:
            continue

    if not text:
        raise ValueError("Could not decode CSV text. File must be valid UTF-8, Windows-1252, or Latin-1.")

    delimiter = detect_csv_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)

    raw_grid = []
    for row in reader:
        cleaned_row = [sanitize_cell_value(c) for c in row]
        raw_grid.append(cleaned_row)

    return raw_grid, detected_encoding, delimiter


def parse_xlsx_grid(file_bytes: bytes, sheet_name: Optional[str] = None) -> Tuple[List[List[Any]], List[str], str]:
    """
    Reads XLSX data cells without formula evaluation using openpyxl.
    Returns: (raw_grid, all_sheet_names, active_sheet_name)
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    all_sheet_names = wb.sheetnames
    if not all_sheet_names:
        raise ValueError("Excel file contains no worksheets.")

    if sheet_name and sheet_name in all_sheet_names:
        ws = wb[sheet_name]
        active_name = sheet_name
    else:
        ws = wb.active
        active_name = ws.title

    raw_grid = []
    for row in ws.iter_rows(values_only=True):
        cleaned_row = [sanitize_cell_value(c) for c in row]
        raw_grid.append(cleaned_row)

    return raw_grid, all_sheet_names, active_name


def extract_single_cell_options(text: str) -> List[str]:
    """
    Extracts options when all options are combined inside a single cell or string.
    Supports formats like:
      A) Apple \n B) Banana \n C) Cherry \n D) Date
      (a) Apple; (b) Banana; (c) Cherry
      1. Apple | 2. Banana | 3. Cherry
    """
    if not text:
        return []

    lines = [l.strip() for l in re.split(r'[\r\n]+', text) if l.strip()]
    if len(lines) >= 2:
        # Check if lines have letter prefixes
        cleaned = [re.sub(r'^(?:\(?[A-Da-d0-9]\)?[\.\)\:\-]\s*|\(?[A-Da-d0-9]\)\s*)', '', l).strip() for l in lines]
        return [c for c in cleaned if c]

    # Check pipe or semicolon separated
    for sep in ('|', ';', '\n'):
        if sep in text:
            parts = [p.strip() for p in text.split(sep) if p.strip()]
            if len(parts) >= 2:
                cleaned = [re.sub(r'^(?:\(?[A-Da-d0-9]\)?[\.\)\:\-]\s*|\(?[A-Da-d0-9]\)\s*)', '', p).strip() for p in parts]
                return [c for c in cleaned if c]

    return []


def analyze_spreadsheet(
    file_bytes: bytes,
    file_format: str,
    custom_mapping: Optional[Dict[str, str]] = None,
    sheet_name: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Full spreadsheet analysis pipeline:
    1. Extracts raw 2D grid.
    2. Detects header row index.
    3. Auto-maps column headers to canonical targets using synonyms.
    4. Applies custom mapping overrides if provided by faculty.
    5. Returns structured rows, proposed mapping, sample preview values, and statistics.
    """
    if file_format == 'csv':
        raw_grid, enc, delim = parse_csv_grid(file_bytes)
        meta = {'encoding': enc, 'delimiter': delim}
        all_sheets = []
        active_sheet = 'Sheet1'
    elif file_format == 'xlsx':
        raw_grid, all_sheets, active_sheet = parse_xlsx_grid(file_bytes, sheet_name=sheet_name)
        meta = {'sheets': all_sheets, 'active_sheet': active_sheet}
    else:
        raise ValueError(f"Unsupported spreadsheet format '{file_format}'.")

    # Filter trailing completely empty rows
    while raw_grid and not any(str(c).strip() for c in raw_grid[-1] if c is not None):
        raw_grid.pop()

    if not raw_grid:
        raise ValueError("Spreadsheet file is completely empty.")

    header_row_idx = detect_header_row(raw_grid)
    raw_headers = [str(c).strip() for c in raw_grid[header_row_idx]]
    data_rows = raw_grid[header_row_idx + 1:]

    # Proposed automatic mapping
    proposed_mapping: Dict[str, str] = {}
    mapping_confidence: Dict[str, float] = {}

    for idx, h in enumerate(raw_headers):
        col_key = f"col_{idx}"
        target, conf = match_column_name(h)
        proposed_mapping[col_key] = target
        mapping_confidence[col_key] = conf

    # Apply custom mapping overrides if provided
    effective_mapping = dict(proposed_mapping)
    if custom_mapping:
        for k, v in custom_mapping.items():
            if v in COLUMN_TARGETS:
                effective_mapping[k] = v

    # Build column metadata with sample values for mapping screen
    column_meta = []
    for idx, h in enumerate(raw_headers):
        col_key = f"col_{idx}"
        samples = []
        for r in data_rows[:5]:
            if idx < len(r) and r[idx] is not None and str(r[idx]).strip():
                samples.append(str(r[idx]).strip()[:80])

        column_meta.append({
            'col_key': col_key,
            'col_index': idx,
            'header': h or f"Column {idx + 1}",
            'proposed_target': proposed_mapping.get(col_key, TARGET_IGNORE),
            'effective_target': effective_mapping.get(col_key, TARGET_IGNORE),
            'confidence': round(mapping_confidence.get(col_key, 0.0), 2),
            'samples': samples,
        })

    # Transform data rows into raw question dicts according to effective mapping
    normalized_records = []
    for row_idx, r in enumerate(data_rows, start=header_row_idx + 2):
        if not r or not any(str(c).strip() for c in r if c is not None):
            continue

        record: Dict[str, Any] = {
            '_row_number': row_idx,
            'options_map': {},
        }

        for idx, cell_val in enumerate(r):
            col_key = f"col_{idx}"
            target = effective_mapping.get(col_key, TARGET_IGNORE)
            if target == TARGET_IGNORE:
                continue

            val_str = str(cell_val).strip() if cell_val is not None else ""
            if not val_str:
                continue

            if target in (TARGET_OPTION_A, TARGET_OPTION_B, TARGET_OPTION_C, TARGET_OPTION_D, TARGET_OPTION_E, TARGET_OPTION_F):
                opt_letter = target.split('_')[1].upper()
                record['options_map'][opt_letter] = val_str
            else:
                record[target] = val_str

        # If options_map has items, assemble options list
        if record['options_map']:
            # Sort options A..F
            sorted_opts = [record['options_map'][k] for k in sorted(record['options_map'].keys())]
            record['options'] = sorted_opts
        elif record.get('options') and isinstance(record['options'], str):
            single_cell_opts = extract_single_cell_options(record['options'])
            if single_cell_opts:
                record['options'] = single_cell_opts

        normalized_records.append(record)

    return {
        'format': file_format,
        'metadata': meta,
        'header_row_index': header_row_idx,
        'total_data_rows': len(normalized_records),
        'columns': column_meta,
        'effective_mapping': effective_mapping,
        'records': normalized_records,
    }

